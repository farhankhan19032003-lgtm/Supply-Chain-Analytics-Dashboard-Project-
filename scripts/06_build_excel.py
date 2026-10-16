#!/usr/bin/env python3
"""
06_build_excel.py  -  builds excel/Supply_Chain_KPI_Workbook.xlsx from the SQL-procedure outputs.
Live formulas: stock cover, stock status (driven by the Config sheet), stock value, KPI tiles, warehouse
status counts, fill-rate trend. Procedure outputs (alerts, RCA, waste, rebalance, assets) are pasted as values
and labelled with the procedure that produced them.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R, D = f"{ROOT}/outputs/sql_results", f"{ROOT}/data/clean"
os.makedirs(f"{ROOT}/excel", exist_ok=True)

F = "Arial"
f_norm, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=16, bold=True, color="1F3864")
f_sub = Font(name=F, size=10, italic=True, color="595959")
f_input = Font(name=F, size=10, color="0000FF")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_tot = PatternFill("solid", fgColor="D9E1F2")
fill_note = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")

wb = Workbook()

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = f_title
    if sub:
        ws["A2"] = sub; ws["A2"].font = f_sub
    ws.sheet_view.showGridLines = False

def fmt_for(col):
    c = col.lower()
    if c.endswith("_inr") or c in ("excess_value_inr",): return "#,##0"
    if c.endswith("_pct") or "pct" in c: return "0.0"
    if c.startswith("weeks_of_cover") or c.startswith("avg_"): return "0.00"
    if c.endswith("_date") or c in ("week_start", "week_end", "order_date", "promised_date"): return "dd-mmm-yy"
    if c.endswith("units") or c.endswith("_qty") or c in ("closing_stock", "demand_units", "units_sold", "lost_units", "stockout_events", "units"): return "#,##0"
    return None

def write_df(ws, df, r0=4, widths=None):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(r0, j, c); cell.font = f_head; cell.fill = fill_head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if pd.isna(v): v = None
            cell = ws.cell(i, j, v); cell.font = f_norm
            nf = fmt_for(df.columns[j - 1])
            if nf and v is not None and not isinstance(v, str): cell.number_format = nf
    for j, c in enumerate(df.columns, 1):
        w = max(len(str(c)) * 0.9, df[c].astype(str).str.len().quantile(0.95) if len(df) else 8) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(w, 11), 46)
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(df.columns))}{r0 + len(df)}"
    return r0 + len(df)

def note(ws, cell, text):
    ws[cell] = text; ws[cell].font = f_sub

# ---------------------------------------------------------------- data
inv = pd.read_csv(f"{R}/tableau_inventory_weekly.csv")
AS_OF = inv["week_end"].max()
snap = inv[inv["week_end"] == AS_OF].sort_values(["warehouse_id", "sku_id"]).reset_index(drop=True)
cfg = pd.read_csv(f"{D}/cfg_thresholds.csv")

# ---------------------------------------------------------------- Cover
ws = wb.active; ws.title = "Cover"
title(ws, "Supply Chain Analytics Platform - KPI Workbook", f"Nationwide (India) distribution network | 12 warehouses | 120 SKUs | 104 weeks | data as of week ending {AS_OF}")
rows = [
    ("Sheet", "What it contains", "Source"),
    ("KPI_Dashboard", "Headline KPIs, status counts by warehouse, two charts (all live formulas)", "Formulas on Stock_Snapshot / Weekly_Trend"),
    ("Config", "Alert thresholds. Change a blue cell and every status in the workbook updates", "cfg_thresholds table"),
    ("Stock_Snapshot", "Latest week, every SKU x warehouse: cover, status, value, margin at risk (formulas)", "vw_stock_status_latest"),
    ("Weekly_Trend", "104-week demand, sales, lost sales, stock value, fill rate", "vw_inventory_enriched"),
    ("Warehouse_Scorecard", "Fill rate, stock-out rate, turns, on-time inbound POs by warehouse", "sp_regional_kpis"),
    ("Supplier_Scorecard", "On-time %, delay, quantity fill, overdue open POs by supplier", "vw_supplier_scorecard"),
    ("Deficit_Alerts", "Every deficit line with inbound-PO risk and weekly margin at risk", "sp_stock_deficit_alerts"),
    ("Overdue_POs", "Open purchase orders past their promise date, with stock exposure", "sp_pipeline_delay_alerts"),
    ("RCA_Stockouts", "Why stock-outs happened: cause summary and warehouse x cause", "sp_stockout_root_cause"),
    ("Excess_Inventory", "Excess / above-target stock, carrying cost, expiry risk", "sp_excess_inventory_waste"),
    ("Rebalance", "Recommended stock transfers from surplus to deficit warehouses", "sp_regional_rebalance"),
    ("Asset_RCA", "Idle, overstretched and ageing assets by warehouse", "sp_asset_waste_rca"),
]
for i, r in enumerate(rows, 4):
    for j, v in enumerate(r, 1):
        c = ws.cell(i, j, v); c.font = f_head if i == 4 else f_norm
        if i == 4: c.fill = fill_head
ws["A19"] = "IMPORTANT - about the data"; ws["A19"].font = f_bold
ws["A20"] = ("All data is SIMULATED (generated by scripts/01_generate_data.py) to mirror a real Indian distribution network. "
             "Supplier names are fictional. Do not present these figures as results from a real employer.")
ws["A20"].font = f_norm; ws["A20"].fill = fill_note; ws["A20"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A20:C21")
ws["A23"] = "Colour code: blue text = editable input; black = formula or pasted result."; ws["A23"].font = f_sub
ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 78; ws.column_dimensions["C"].width = 40

# ---------------------------------------------------------------- Config
wc = wb.create_sheet("Config")
title(wc, "Alert thresholds", "Blue cells are inputs. Stock_Snapshot status formulas read rows 5-7.")
for j, h in enumerate(["Parameter", "Value", "Description"], 1):
    c = wc.cell(4, j, h); c.font = f_head; c.fill = fill_head
for i, r in enumerate(cfg.itertuples(index=False), 5):
    wc.cell(i, 1, r.param).font = f_norm
    c = wc.cell(i, 2, float(r.param_value)); c.font = f_input
    wc.cell(i, 3, r.description).font = f_norm
wc.column_dimensions["A"].width = 38; wc.column_dimensions["B"].width = 12; wc.column_dimensions["C"].width = 80
note(wc, f"A{6 + len(cfg)}", "Source: project assumptions (typical FMCG DC cover norms); holding cost 22% p.a. is an assumed carrying-cost rate.")

# ---------------------------------------------------------------- Stock_Snapshot
ss = wb.create_sheet("Stock_Snapshot")
title(ss, "Stock snapshot - latest week", f"Week ending {AS_OF}. Cover, status, value and margin-at-risk are formulas driven by the Config sheet.")
cols = ["warehouse_id", "warehouse_name", "region", "sku_id", "sku_name", "category", "abc_class", "closing_stock",
        "avg_demand_4w", "weeks_of_cover", "stock_status", "unit_cost_inr", "unit_price_inr", "stock_value_inr",
        "weekly_margin_at_risk_inr", "on_order_units"]
for j, h in enumerate(cols, 1):
    c = ss.cell(4, j, h); c.font = f_head; c.fill = fill_head; c.alignment = Alignment(horizontal="center", wrap_text=True)
ss.row_dimensions[4].height = 30
r0 = 5
for i, r in snap.iterrows():
    n = r0 + i
    vals = [r.warehouse_id, r.warehouse_name, r.region, r.sku_id, r.sku_name, r.category, r.abc_class,
            int(r.closing_stock), round(float(r.avg_demand_4w), 2)]
    for j, v in enumerate(vals, 1): ss.cell(n, j, v).font = f_norm
    ss.cell(n, 10, f'=IF(I{n}>0,H{n}/I{n},"")')
    ss.cell(n, 11, f'=IF(I{n}<=0,"NO DEMAND",IF(H{n}=0,"STOCKOUT",IF(J{n}<Config!$B$5,"CRITICAL",IF(J{n}<Config!$B$6,"LOW",IF(J{n}>Config!$B$7,"EXCESS","HEALTHY")))))')
    ss.cell(n, 12, float(r.unit_cost_inr)); ss.cell(n, 13, float(r.unit_price_inr))
    ss.cell(n, 14, f"=H{n}*L{n}")
    ss.cell(n, 15, f"=IF(I{n}>0,I{n}*(M{n}-L{n}),0)")
    ss.cell(n, 16, int(r.on_order_units))
    for j in range(10, 17): ss.cell(n, j).font = f_norm
    ss.cell(n, 8).number_format = "#,##0"; ss.cell(n, 10).number_format = "0.00"
    for j in (12, 13, 14, 15): ss.cell(n, j).number_format = "#,##0"
    ss.cell(n, 16).number_format = "#,##0"
LAST = r0 + len(snap) - 1
for w, col in zip([10, 16, 9, 9, 26, 24, 7, 11, 11, 10, 12, 11, 11, 13, 14, 11], range(1, 17)):
    ss.column_dimensions[get_column_letter(col)].width = w
ss.freeze_panes = "E5"; ss.auto_filter.ref = f"A4:P{LAST}"
rng = f"K5:K{LAST}"
for txt, col in [("STOCKOUT", "F8CBAD"), ("CRITICAL", "F4B084"), ("LOW", "FFE699"), ("EXCESS", "BDD7EE"), ("HEALTHY", "C6E0B4")]:
    ss.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", bgColor=col, fgColor=col)))

# ---------------------------------------------------------------- Weekly_Trend
wt = wb.create_sheet("Weekly_Trend")
title(wt, "Weekly network trend", "Aggregated from vw_inventory_enriched. Fill rate is a formula.")
g = inv.groupby("week_start").agg(demand_units=("demand_units", "sum"), units_sold=("units_sold", "sum"),
                                  lost_units=("lost_units", "sum"), stock_value_inr=("stock_value_inr", "sum"),
                                  lost_revenue_inr=("lost_revenue_inr", "sum")).reset_index()
g["week_start"] = pd.to_datetime(g["week_start"])
last_t = write_df(wt, g, 4)
wt.cell(4, 7, "fill_rate").font = f_head; wt.cell(4, 7).fill = fill_head
for n in range(5, last_t + 1):
    c = wt.cell(n, 7, f"=IF(B{n}>0,C{n}/B{n},0)"); c.number_format = "0.0%"; c.font = f_norm
    wt.cell(n, 1).number_format = "dd-mmm-yy"
wt.column_dimensions["G"].width = 11
wt.auto_filter.ref = f"A4:G{last_t}"
T0, T1 = 5, last_t

# ---------------------------------------------------------------- Warehouse_Scorecard
k = pd.read_csv(f"{R}/kpi_by_warehouse.csv").merge(
    pd.read_csv(f"{R}/po_perf_by_warehouse.csv")[["warehouse_id", "on_time_pct", "avg_days_vs_promise", "qty_fill_pct"]].rename(
        columns={"on_time_pct": "inbound_on_time_pct", "avg_days_vs_promise": "inbound_avg_days_vs_promise", "qty_fill_pct": "inbound_qty_fill_pct"}),
    on="warehouse_id")
wsw = wb.create_sheet("Warehouse_Scorecard")
title(wsw, "Warehouse scorecard - 104 weeks", "Source: sp_regional_kpis. Totals row and lost-revenue share are formulas.")
lastw = write_df(wsw, k, 4)
tr = lastw + 1
wsw.cell(tr, 1, "NETWORK").font = f_bold
for j, c in enumerate(k.columns, 1):
    L = get_column_letter(j)
    if c in ("demand_units", "units_sold", "lost_units", "avg_inventory_value_inr", "lost_revenue_inr", "lost_margin_inr"):
        cell = wsw.cell(tr, j, f"=SUM({L}5:{L}{lastw})"); cell.number_format = "#,##0"
    elif c == "fill_rate_pct":
        dj, sj = get_column_letter(list(k.columns).index("demand_units") + 1), get_column_letter(list(k.columns).index("units_sold") + 1)
        cell = wsw.cell(tr, j, f"=100*{sj}{tr}/{dj}{tr}"); cell.number_format = "0.00"
    else:
        continue
    cell.font = f_bold
for j in range(1, len(k.columns) + 1): wsw.cell(tr, j).fill = fill_tot
sc = len(k.columns) + 1
wsw.cell(4, sc, "lost_revenue_share_pct").font = f_head; wsw.cell(4, sc).fill = fill_head
LRc = get_column_letter(list(k.columns).index("lost_revenue_inr") + 1)
for n in range(5, lastw + 1):
    c = wsw.cell(n, sc, f"=100*{LRc}{n}/{LRc}${tr}"); c.number_format = "0.0"; c.font = f_norm
wsw.column_dimensions[get_column_letter(sc)].width = 14

# ---------------------------------------------------------------- Supplier_Scorecard
sup = pd.read_csv(f"{R}/supplier_scorecard.csv")
wsu = wb.create_sheet("Supplier_Scorecard")
title(wsu, "Supplier scorecard", "Source: vw_supplier_scorecard (sorted worst on-time first). Red = on-time below 60%.")
lasts = write_df(wsu, sup, 4)
oc = get_column_letter(list(sup.columns).index("on_time_pct") + 1)
wsu.conditional_formatting.add(f"{oc}5:{oc}{lasts}", CellIsRule(operator="lessThan", formula=["60"], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))

# ---------------------------------------------------------------- alert sheets & other procedure outputs
sheets = [
    ("Deficit_Alerts", "Stock-deficit alerts", "Source: sp_stock_deficit_alerts. Sorted STOCKOUT > CRITICAL > LOW, then margin at risk.", pd.read_csv(f"{R}/stock_deficit_alerts.csv")),
    ("Overdue_POs", "Pipeline delays - open POs past promise date", "Source: sp_pipeline_delay_alerts (>= warn threshold). Most overdue first.", pd.read_csv(f"{R}/pipeline_delay_alerts.csv")),
]
for name, t, s, df in sheets:
    w = wb.create_sheet(name); title(w, t, s); write_df(w, df, 4)

wr = wb.create_sheet("RCA_Stockouts")
title(wr, "Root-cause analysis of stock-out events", "Source: sp_stockout_root_cause. One primary cause per event; precedence documented in the procedure header.")
cs = pd.read_csv(f"{R}/rca_cause_summary.csv")
end1 = write_df(wr, cs, 4)
wr.freeze_panes = None; wr.auto_filter.ref = None
wr.cell(end1 + 1, 1, "TOTAL").font = f_bold
for col in ("stockout_events", "lost_units", "lost_revenue_inr", "lost_margin_inr", "pct_of_lost_revenue"):
    j = list(cs.columns).index(col) + 1; L = get_column_letter(j)
    c = wr.cell(end1 + 1, j, f"=SUM({L}5:{L}{end1})"); c.font = f_bold; c.number_format = "#,##0" if col != "pct_of_lost_revenue" else "0.0"
r2 = end1 + 4
wr.cell(r2 - 1, 1, "Lost revenue (INR) by warehouse and primary cause").font = f_bold
mat = pd.read_csv(f"{R}/rca_warehouse_matrix.csv")
end2 = write_df(wr, mat, r2); wr.freeze_panes = None; wr.auto_filter.ref = None
wr.cell(end2 + 1, 1, "TOTAL").font = f_bold
for col in mat.columns[3:]:
    j = list(mat.columns).index(col) + 1; L = get_column_letter(j)
    c = wr.cell(end2 + 1, j, f"=SUM({L}{r2 + 1}:{L}{end2})"); c.font = f_bold; c.number_format = "#,##0"
ov = pd.read_csv(f"{R}/rca_factor_overlap.csv")
wr.cell(end2 + 4, 1, "How often each factor was present (events can have several)").font = f_bold
write_df(wr, ov, end2 + 5); wr.freeze_panes = None; wr.auto_filter.ref = None
wr.column_dimensions["A"].width = 46

we = wb.create_sheet("Excess_Inventory")
title(we, "Excess and above-target inventory", "Source: sp_excess_inventory_waste. Above-target = stock beyond 4 weeks of cover; excess = beyond 8 weeks.")
ew = pd.read_csv(f"{R}/excess_by_warehouse.csv")
e1 = write_df(we, ew, 4); we.freeze_panes = None; we.auto_filter.ref = None
we.cell(e1 + 1, 1, "NETWORK").font = f_bold
for col in ew.columns[3:]:
    j = list(ew.columns).index(col) + 1; L = get_column_letter(j)
    if "pct" in col: continue
    c = we.cell(e1 + 1, j, f"=SUM({L}5:{L}{e1})"); c.font = f_bold; c.number_format = "#,##0"
we.cell(e1 + 4, 1, "By category").font = f_bold
write_df(we, pd.read_csv(f"{R}/excess_by_category.csv"), e1 + 5); we.freeze_panes = None; we.auto_filter.ref = None
we.column_dimensions["A"].width = 24

wrb = wb.create_sheet("Rebalance")
title(wrb, "Recommended stock transfers", "Source: sp_regional_rebalance. Indicative, one replenishment cycle; freight rates are assumptions (Config).")
tf = pd.read_csv(f"{R}/rebalance_transfers.csv")
lr_ = write_df(wrb, tf, 4)
wrb.cell(lr_ + 1, 1, "TOTAL").font = f_bold
for col in ("transfer_units", "stock_value_moved_inr", "est_freight_inr", "margin_protected_inr", "est_net_benefit_inr"):
    j = list(tf.columns).index(col) + 1; L = get_column_letter(j)
    c = wrb.cell(lr_ + 1, j, f"=SUM({L}5:{L}{lr_})"); c.font = f_bold; c.number_format = "#,##0"

wa = wb.create_sheet("Asset_RCA")
title(wa, "Asset waste root-cause analysis (Oct-2024 to Sep-2026)", "Source: sp_asset_waste_rca. surplus_assets > 0 = over-provisioned; < 0 = shortfall.")
aw = pd.read_csv(f"{R}/asset_rca_warehouse.csv")
a1 = write_df(wa, aw, 4); wa.freeze_panes = None; wa.auto_filter.ref = None
wa.cell(a1 + 1, 1, "NETWORK").font = f_bold
for col in ("assets", "breakdowns", "maintenance_cost_inr", "idle_cost_inr", "net_surplus_assets"):
    j = list(aw.columns).index(col) + 1; L = get_column_letter(j)
    c = wa.cell(a1 + 1, j, f"=SUM({L}5:{L}{a1})"); c.font = f_bold; c.number_format = "#,##0"
wa.cell(a1 + 4, 1, "Detail by warehouse and asset type").font = f_bold
ad = pd.read_csv(f"{R}/asset_rca_detail.csv")
write_df(wa, ad, a1 + 5)
wa.column_dimensions["A"].width = 12

# ---------------------------------------------------------------- KPI_Dashboard (formulas)
kd = wb.create_sheet("KPI_Dashboard", 1)
title(kd, "Network KPI dashboard", f"Live formulas on Stock_Snapshot, Weekly_Trend and the scorecards | as of {AS_OF}")
SS = lambda col: f"Stock_Snapshot!${col}$5:${col}${LAST}"
WT = lambda col: f"Weekly_Trend!${col}${T0}:${col}${T1}"
items = [
    ("Network stock value (INR)", f"=SUM({SS('N')})", "#,##0", "Closing stock x unit cost, latest week"),
    ("SKU x warehouse lines tracked", f"=COUNTA({SS('D')})", "#,##0", ""),
    ("Lines in STOCKOUT", f'=COUNTIF({SS("K")},"STOCKOUT")', "#,##0", "closing stock = 0 with demand"),
    ("Lines CRITICAL", f'=COUNTIF({SS("K")},"CRITICAL")', "#,##0", "cover below Config critical threshold"),
    ("Lines LOW", f'=COUNTIF({SS("K")},"LOW")', "#,##0", "cover below Config low threshold"),
    ("Lines EXCESS", f'=COUNTIF({SS("K")},"EXCESS")', "#,##0", "cover above Config excess threshold"),
    ("Deficit lines as % of all lines", "=(B7+B8+B9)/B6", "0.0%", "STOCKOUT + CRITICAL + LOW"),
    ("Weekly margin at risk on deficit lines (INR)", f'=SUMIFS({SS("O")},{SS("K")},"STOCKOUT")+SUMIFS({SS("O")},{SS("K")},"CRITICAL")+SUMIFS({SS("O")},{SS("K")},"LOW")', "#,##0", "avg weekly demand x unit margin"),
    ("Network fill rate, 104 weeks", f"=SUM({WT('C')})/SUM({WT('B')})", "0.00%", "units sold / units demanded"),
    ("Revenue lost to stock-outs, 104 weeks (INR)", f"=SUM({WT('F')})", "#,##0", "lost units x selling price"),
    ("Overdue open POs (>= warn days)", f"=COUNTA(Overdue_POs!$A$5:$A${4 + len(sheets[1][3])})", "#,##0", "pipeline-delay alerts"),
    ("Value of overdue open POs (INR)", f"=SUM(Overdue_POs!$O$5:$O${4 + len(sheets[1][3])})", "#,##0", "PO value of the above"),
]
for j, h in enumerate(["Metric", "Value", "Definition"], 1):
    c = kd.cell(4, j, h); c.font = f_head; c.fill = fill_head
for i, (m, f, nf, d) in enumerate(items, 5):
    kd.cell(i, 1, m).font = f_norm
    c = kd.cell(i, 2, f); c.number_format = nf; c.font = f_bold
    kd.cell(i, 3, d).font = f_sub
# sanity: PO value column letter
po_cols = list(sheets[1][3].columns)
po_val_col = get_column_letter(po_cols.index("po_value_inr") + 1)
kd["B16"] = f"=SUM(Overdue_POs!${po_val_col}$5:${po_val_col}${4 + len(sheets[1][3])})"

hr = 19
heads = ["warehouse_id", "warehouse_name", "region", "stock_value_inr", "pct_of_network_stock", "STOCKOUT", "CRITICAL", "LOW", "EXCESS", "deficit_lines_pct", "fill_rate_pct (104 wks)"]
for j, h in enumerate(heads, 1):
    c = kd.cell(hr, j, h); c.font = f_head; c.fill = fill_head; c.alignment = Alignment(wrap_text=True, horizontal="center")
kd.row_dimensions[hr].height = 30
whs = snap[["warehouse_id", "warehouse_name", "region"]].drop_duplicates().reset_index(drop=True)
for i, r in whs.iterrows():
    n = hr + 1 + i
    kd.cell(n, 1, r.warehouse_id); kd.cell(n, 2, r.warehouse_name); kd.cell(n, 3, r.region)
    kd.cell(n, 4, f"=SUMIFS({SS('N')},{SS('A')},$A{n})").number_format = "#,##0"
    kd.cell(n, 5, f"=D{n}/$B$5").number_format = "0.0%"
    for j, st in zip((6, 7, 8, 9), ("STOCKOUT", "CRITICAL", "LOW", "EXCESS")):
        kd.cell(n, j, f'=COUNTIFS({SS("A")},$A{n},{SS("K")},"{st}")').number_format = "#,##0"
    kd.cell(n, 10, f"=(F{n}+G{n}+H{n})/COUNTIFS({SS('A')},$A{n})").number_format = "0.0%"
    kd.cell(n, 11, f"=INDEX(Warehouse_Scorecard!$G$5:$G$16,MATCH($A{n},Warehouse_Scorecard!$A$5:$A$16,0))").number_format = "0.00"
    for j in range(1, 12): kd.cell(n, j).font = f_norm
wl = hr + len(whs)
kd.cell(wl + 1, 1, "NETWORK").font = f_bold
for j, L in ((4, "D"), (6, "F"), (7, "G"), (8, "H"), (9, "I")):
    c = kd.cell(wl + 1, j, f"=SUM({L}{hr + 1}:{L}{wl})"); c.font = f_bold; c.number_format = "#,##0"
for j in range(1, 12): kd.cell(wl + 1, j).fill = fill_tot
for col, w in zip("ABCDEFGHIJK", [44, 18, 40, 16, 12, 11, 11, 9, 9, 12, 14]):
    kd.column_dimensions[col].width = w
kd.conditional_formatting.add(f"J{hr + 1}:J{wl}", CellIsRule(operator="greaterThan", formula=["0.3"], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))
kd.conditional_formatting.add(f"E{hr + 1}:E{wl}", CellIsRule(operator="greaterThan", formula=["0.2"], fill=PatternFill("solid", bgColor="BDD7EE", fgColor="BDD7EE")))
note(kd, f"A{wl + 3}", "Red = more than 30% of the warehouse's SKU lines are in deficit.  Blue = holds more than 20% of all network stock value.")

ch1 = BarChart(); ch1.type = "col"; ch1.title = "Stock value by warehouse (INR)"; ch1.style = 10
ch1.add_data(Reference(kd, min_col=4, min_row=hr, max_row=wl), titles_from_data=True)
ch1.set_categories(Reference(kd, min_col=2, min_row=hr + 1, max_row=wl))
ch1.legend = None; ch1.height, ch1.width = 8, 16
ch1.y_axis.numFmt = "#,##0"; ch1.y_axis.delete = False; ch1.x_axis.delete = False
kd.add_chart(ch1, f"A{wl + 5}")

ch2 = LineChart(); ch2.title = "Weekly network fill rate"; ch2.style = 12
ch2.add_data(Reference(wt, min_col=7, min_row=4, max_row=T1), titles_from_data=True)
ch2.set_categories(Reference(wt, min_col=1, min_row=T0, max_row=T1))
ch2.legend = None; ch2.height, ch2.width = 8, 16
ch2.y_axis.numFmt = "0%"; ch2.y_axis.scaling.min = 0.9; ch2.y_axis.scaling.max = 1.0
ch2.x_axis.number_format = "mmm-yy"; ch2.y_axis.delete = False; ch2.x_axis.delete = False
kd.add_chart(ch2, f"D{wl + 5}")

out = f"{ROOT}/excel/Supply_Chain_KPI_Workbook.xlsx"
wb.save(out)
print("saved", out)
